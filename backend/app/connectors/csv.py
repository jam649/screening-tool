"""Bounded CSV/XLSX previews and explicit mapped-row validation."""

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.connectors.base import PartyInput, RejectedRow, SyncBatch
from app.screening.normalize import normalize_country, normalize_name, party_record_hash

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_XLSX_BYTES = 64 * 1024 * 1024
MAX_ROWS = 25_000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 10_000
MAX_REJECTED_DETAILS = 500
MAPPABLE_FIELDS = frozenset(
    {
        "name",
        "party_type",
        "role",
        "alt_names",
        "address_lines",
        "city",
        "region",
        "postal_code",
        "country",
        "external_id",
    }
)
STABLE_ID_WARNING = (
    "Map a stable customer or vendor ID to keep changes in the same party file. "
    "Without an ID, changed screening fields create a new party file; previous history is retained."
)


class UploadValidationError(ValueError):
    pass


@dataclass(frozen=True)
class UploadTable:
    headers: list[str]
    rows: list[dict[str, str]]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        raise UploadValidationError("An upload cell exceeds the 10,000-character limit")
    if "\x00" in text:
        raise UploadValidationError("Upload cells cannot contain NUL characters")
    return text


def read_upload(data: bytes, filename: str) -> UploadTable:
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise UploadValidationError("Upload must be nonempty and no larger than 10 MiB")
    extension = Path(filename).suffix.casefold()
    if extension == ".csv":
        matrix: list[list[str]] = []
        try:
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            decoded = data.decode(encoding)
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(decoded, newline=""), dialect=dialect, strict=True)
            matrix = []
            for row in reader:
                if len(matrix) > MAX_ROWS:
                    raise UploadValidationError("Upload exceeds 25,000 rows")
                if len(row) > MAX_COLUMNS:
                    raise UploadValidationError("Upload exceeds 100 columns")
                matrix.append([_text(value) for value in row])
        except (UnicodeError, csv.Error) as exc:
            raise UploadValidationError(
                "CSV must be valid UTF-8 or BOM-marked UTF-16 text"
            ) from exc
    elif extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > 1000
                    or sum(member.file_size for member in members) > MAX_EXPANDED_XLSX_BYTES
                ):
                    raise UploadValidationError("Expanded spreadsheet exceeds safe size limits")
                if any(member.flag_bits & 1 for member in members):
                    raise UploadValidationError("Encrypted spreadsheets are not supported")
                if any("vbaProject" in member.filename for member in members):
                    raise UploadValidationError("Spreadsheet macros are not supported")
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            try:
                worksheet = workbook.worksheets[0]
                # Ignore author-supplied dimensions; streaming row/cell caps are authoritative.
                worksheet.reset_dimensions()
                matrix = []
                for row in worksheet.iter_rows(values_only=True):
                    if len(matrix) > MAX_ROWS:
                        raise UploadValidationError("Upload exceeds 25,000 rows")
                    if len(row) > MAX_COLUMNS:
                        raise UploadValidationError("Upload exceeds 100 columns")
                    matrix.append([_text(value) for value in row])
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            if isinstance(exc, UploadValidationError):
                raise
            raise UploadValidationError("Spreadsheet is not a readable XLSX file") from exc
    else:
        raise UploadValidationError("Only .csv and .xlsx uploads are supported")
    if not matrix:
        raise UploadValidationError("Upload does not contain a header row")
    headers = [value.strip() for value in matrix[0]]
    # Ignore trailing blank header cells only when every row is blank in those columns.
    while (
        headers
        and not headers[-1]
        and all(len(row) < len(headers) or not row[len(headers) - 1] for row in matrix[1:])
    ):
        headers.pop()
    if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
        raise UploadValidationError("Column headings must be nonempty and unique")
    rows = []
    for row in matrix[1:]:
        if any(row[len(headers) :]):
            raise UploadValidationError("A data row contains more cells than the header")
        values = row[: len(headers)] + [""] * max(0, len(headers) - len(row))
        rows.append(dict(zip(headers, values, strict=True)))
    return UploadTable(headers, rows)


def preview_upload(data: bytes, filename: str) -> dict[str, Any]:
    table = read_upload(data, filename)
    return {
        "headers": table.headers,
        "rows": table.rows[:20],
        "total_rows": len(table.rows),
        "stable_id_warning": STABLE_ID_WARNING,
    }


class CSVConnector:
    def __init__(self, data: bytes, filename: str, mapping: dict[str, str | list[str]]):
        self.table = read_upload(data, filename)
        self.mapping = mapping
        if "name" not in mapping or not isinstance(mapping["name"], str):
            raise UploadValidationError("Map a name column before importing")
        for field, columns in mapping.items():
            if field not in MAPPABLE_FIELDS:
                raise UploadValidationError(f"Unsupported mapped field: {field}")
            if not isinstance(columns, (str, list)) or (
                isinstance(columns, list) and field not in {"address_lines", "alt_names"}
            ):
                raise UploadValidationError("Only addresses and aliases accept multiple columns")
            selected = [columns] if isinstance(columns, str) else columns
            if not selected or any(
                not isinstance(column, str) or column not in self.table.headers
                for column in selected
            ):
                raise UploadValidationError("Mapped columns must exist in the upload")

    def normalize(self, raw: dict[str, Any]) -> PartyInput:
        def value(field: str) -> str:
            column = self.mapping.get(field)
            return str(raw.get(column, "")).strip() if isinstance(column, str) else ""

        def values(field: str) -> list[str]:
            columns = self.mapping.get(field, [])
            columns = [columns] if isinstance(columns, str) else columns
            return [str(raw[column]).strip() for column in columns if str(raw[column]).strip()]

        name = value("name")
        if not name or not normalize_name(name).tokens:
            raise ValueError("Name is required and must contain letters or numbers")
        party_type = value("party_type").casefold() or "company"
        party_type = {"entity": "company", "person": "individual"}.get(party_type, party_type)
        role = value("role").casefold() or "customer"
        if party_type not in {"company", "individual"}:
            raise ValueError("Party type must be company or individual")
        if role not in {"customer", "vendor", "both"}:
            raise ValueError("Role must be customer, vendor, or both")
        country = normalize_country(value("country"), strict=True)
        fields: dict[str, Any] = {
            "name": name,
            "party_type": party_type,
            "role": role,
            "alt_names": values("alt_names"),
            "address_lines": values("address_lines"),
            "city": value("city") or None,
            "region": value("region") or None,
            "postal_code": value("postal_code") or None,
            "country_iso2": country,
        }
        record_hash = party_record_hash(fields)
        external_id = value("external_id") or f"generated:{record_hash}"
        if len(external_id) > 500:
            raise ValueError("External ID exceeds 500 characters")
        return PartyInput(external_id=external_id, **fields, raw=raw, record_hash=record_hash)

    def full_sync(self) -> SyncBatch:
        parties = []
        rejected: list[RejectedRow] = []
        rejected_count = 0
        seen = set()
        for number, raw in enumerate(self.table.rows, start=2):
            try:
                party = self.normalize(raw)
                if party.external_id in seen:
                    raise ValueError("Duplicate external ID; map a unique stable ID for each row")
                seen.add(party.external_id)
                parties.append(party)
            except ValueError as exc:
                rejected_count += 1
                if len(rejected) < MAX_REJECTED_DETAILS:
                    rejected.append(RejectedRow(number, str(exc)))
        return SyncBatch(parties, rejected, complete=True, rejected_count=rejected_count)

    def incremental_sync(self, cursor: str | None) -> SyncBatch:
        # Uploads represent a full source snapshot, never an incremental deletion signal.
        return self.full_sync()
